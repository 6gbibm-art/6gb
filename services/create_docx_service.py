import io
from docx import Document
from fastapi.responses import StreamingResponse
def create_docx(letter: str) -> StreamingResponse:
    document = Document()
    document.add_heading("Cover Letter",level=1)
    for paragraph in letter.split("\n"):
        document.add_paragraph(paragraph)
    buffer = io.BytesIO()
    document.save(buffer)
    buffer.seek(0)
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={
            "Content-Disposition":
            'attachment; filename="Cover_Letter.docx"'
        }
    )